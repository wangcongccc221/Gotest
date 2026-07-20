from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "outputs" / "一条数据的旅程-前端到FSM再回到界面.pptx"

EXPECTED_TITLES = (
    "一条数据的旅程",
    "三个角色，一条链路",
    "完整链路：连接、请求、指令、返回、显示",
    "连接成功，只代表通道已经建立",
    "前端发的是 JSON，不是 SYNC",
    "Go 按 type 路由请求",
    "Go 把 JSON 翻译成 DISPLAY_ON",
    "FSM 从 1128 返回，Go 分三段读取",
    "配置即时转发，统计先缓存再定时广播",
    "Go 每秒取快照，再广播给所有连接",
    "前端按数据类别字段分流，状态变化驱动显示",
    "八个节点定位：数据停在哪里",
)

REQUIRED_TERMS = {
    1: ("前端请求", "FSM 返回", "页面显示"),
    2: ("ArkTS 前端", "Go 后端", "FSM", "协议翻译"),
    3: ("requestStGlobal", "DISPLAY_ON", "1128", "数据类别字段", "KPI 状态"),
    4: ("openCallback", "requestHomeStats", "不等待 ready"),
    5: ("requestStGlobal", "requestHomeStats", "requestStatistics", "聚合缓存", "统计缓存"),
    6: ("handleIncoming", "handleRequestStGlobal", "goroutine"),
    7: ("SYNC", "0x0019", "1279", "16 字节"),
    8: ("1128", "recvCTCPSync", "recvCTCPCommand", "payload"),
    9: ("ParseData", "cmdFSMConfig", "cmdFSMStatistics", "缓存"),
    10: ("1 秒", "2.5 秒", "所有连接", "topic"),
    11: ("数据类别字段", "AppStorage", "@StorageLink", "homeStats"),
    12: ("排查", "复盘", "WebSocket open", "FSM -> 1128"),
}

FORBIDDEN_PHRASES = (
    "主题订阅",
    "订阅 topic",
    "发布/订阅",
    "mmap",
)

MIN_FONT_BY_PREFIX = {
    "title:": 28.0,
    "body:": 18.0,
    "code:": 14.0,
    "label:": 11.0,
    "small:": 9.0,
    "footer:": 9.0,
}


def slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def minimum_font_for_shape(shape_name: str) -> float:
    for prefix, minimum in MIN_FONT_BY_PREFIX.items():
        if shape_name.startswith(prefix):
            return minimum
    return 10.0


def verify_text_fonts(slide, slide_number: int) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
            continue
        minimum = minimum_font_for_shape(shape.name)
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                assert run.font.size is not None, (
                    f"slide {slide_number} shape {shape.name!r} has inherited font size"
                )
                assert run.font.size.pt + 0.01 >= minimum, (
                    f"slide {slide_number} shape {shape.name!r} uses "
                    f"{run.font.size.pt:.1f}pt; expected at least {minimum:.1f}pt"
                )
                if shape.name.startswith("code:"):
                    assert run.font.name == "Consolas", (
                        f"slide {slide_number} code uses {run.font.name!r}, not Consolas"
                    )


def verify_shape_bounds(prs: Presentation, slide, slide_number: int) -> None:
    tolerance = 1_000
    for shape in slide.shapes:
        left = int(shape.left)
        top = int(shape.top)
        right = left + int(shape.width)
        bottom = top + int(shape.height)
        assert left >= -tolerance and top >= -tolerance, (
            f"slide {slide_number} shape {shape.name!r} starts outside the slide"
        )
        assert right <= int(prs.slide_width) + tolerance, (
            f"slide {slide_number} shape {shape.name!r} exceeds right edge"
        )
        assert bottom <= int(prs.slide_height) + tolerance, (
            f"slide {slide_number} shape {shape.name!r} exceeds bottom edge"
        )


def rectangles_overlap(first, second) -> bool:
    left = max(int(first.left), int(second.left))
    top = max(int(first.top), int(second.top))
    right = min(int(first.left + first.width), int(second.left + second.width))
    bottom = min(int(first.top + first.height), int(second.top + second.height))
    if right <= left or bottom <= top:
        return False
    # Ignore hairline contact and only report a meaningful text-on-text collision.
    return (right - left) * (bottom - top) > 20_000_000


def verify_text_overlaps(slide, slide_number: int) -> None:
    text_shapes = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip()
        and not shape.name.startswith("decor:")
    ]
    for index, first in enumerate(text_shapes):
        for second in text_shapes[index + 1 :]:
            assert not rectangles_overlap(first, second), (
                f"slide {slide_number} text boxes overlap: "
                f"{first.name!r} and {second.name!r}"
            )


def verify_code_density(slide, slide_number: int) -> None:
    if slide_number < 4 or slide_number > 11:
        return
    code_shapes = [
        shape
        for shape in slide.shapes
        if shape.name.startswith("code:") and getattr(shape, "has_text_frame", False)
    ]
    assert code_shapes, f"slide {slide_number} has no code excerpt"
    code_lines = sum(
        sum(1 for line in shape.text.splitlines() if line.strip())
        for shape in code_shapes
    )
    assert 8 <= code_lines <= 24, (
        f"slide {slide_number} has {code_lines} code lines; expected a focused excerpt"
    )


def verify_visual_primitives(slide, slide_number: int) -> None:
    pictures = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert not pictures, f"slide {slide_number} contains an unexpected picture"


def verify() -> None:
    assert DECK.exists(), f"missing deck: {DECK}"
    assert DECK.stat().st_size > 20_000, "deck is unexpectedly small or empty"

    with ZipFile(DECK) as archive:
        assert archive.testzip() is None, "PPTX package contains a corrupt member"
        assert "ppt/presentation.xml" in archive.namelist()
        presentation_xml = archive.read("ppt/presentation.xml").decode("utf-8")
        assert "gradFill" not in presentation_xml

    prs = Presentation(DECK)
    assert len(prs.slides) == 12, f"expected 12 slides, got {len(prs.slides)}"
    assert abs((prs.slide_width / prs.slide_height) - (16 / 9)) < 0.01

    all_deck_text: list[str] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        text = slide_text(slide)
        all_deck_text.append(text)
        assert EXPECTED_TITLES[slide_number - 1] in text, (
            f"slide {slide_number} has the wrong or missing title"
        )
        assert f"{slide_number:02d} / 12" in text, (
            f"slide {slide_number} is missing its page number"
        )
        for term in REQUIRED_TERMS[slide_number]:
            assert term in text, f"slide {slide_number} is missing {term!r}"
        if 4 <= slide_number <= 5 or slide_number == 11:
            assert "E:/new/my_harmony" in text, (
                f"slide {slide_number} is missing the real ArkTS source root"
            )
        if 6 <= slide_number <= 10:
            assert "E:/goTest" in text, (
                f"slide {slide_number} is missing the real Go source root"
            )
        verify_shape_bounds(prs, slide, slide_number)
        verify_text_fonts(slide, slide_number)
        verify_text_overlaps(slide, slide_number)
        verify_code_density(slide, slide_number)
        verify_visual_primitives(slide, slide_number)

    joined = "\n".join(all_deck_text)
    for phrase in FORBIDDEN_PHRASES:
        assert phrase not in joined, f"deck contains forbidden phrase {phrase!r}"

    print(
        "PASS: 12-slide DISPLAY_ON training deck verified "
        f"({DECK.stat().st_size:,} bytes)"
    )


if __name__ == "__main__":
    verify()
