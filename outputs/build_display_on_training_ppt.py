from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable, Sequence
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "一条数据的旅程-前端到FSM再回到界面.pptx"
GO_ROOT = ROOT / "go" / "ohos" / "Tcp"
FRONTEND_ROOT = Path("E:/new/my_harmony")

COLORS = {
    "paper": "F7F8FA",
    "white": "FFFFFF",
    "ink": "17202A",
    "muted": "5F6B76",
    "line": "D9DEE5",
    "request": "2563A6",
    "request_light": "EAF2FA",
    "response": "2D7D5A",
    "response_light": "EAF5EF",
    "warning": "B46620",
    "warning_light": "F9EFE5",
    "backend": "52606D",
    "backend_light": "EEF1F4",
    "code": "18222D",
    "code_muted": "AFC0CD",
    "soft": "EEF1F4",
}

FONT_UI = "Microsoft YaHei"
FONT_CODE = "Consolas"

SLIDE_W = 13.333333
SLIDE_H = 7.5
CONTENT_X = 0.62
CONTENT_W = 12.09


def source_excerpt(
    path: Path,
    ranges: Sequence[tuple[int, int]],
    *,
    max_chars: int = 64,
) -> str:
    """Return verbatim source lines with explicit omission markers."""
    lines = path.read_text(encoding="utf-8").splitlines()
    output: list[str] = []
    for index, (start, end) in enumerate(ranges):
        if index > 0:
            output.append("...")
        if start < 1 or end < start or end > len(lines):
            raise ValueError(f"invalid source range {path}:{start}-{end}")
        for line_number in range(start, end + 1):
            rendered = f"{line_number:04d} {lines[line_number - 1].expandtabs(4)}"
            if len(rendered) > max_chars:
                rendered = rendered[: max_chars - 4].rstrip() + " ..."
            output.append(rendered)
    return "\n".join(output)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_shape_name(shape, name: str) -> None:
    shape._element.nvSpPr.cNvPr.set("name", name)


def style_fill(shape, color: str, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def style_line(shape, color: str, width: float = 1.0, transparency: int = 0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str,
    line: str | None = None,
    line_width: float = 1.0,
    name: str = "decor:rect",
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    set_shape_name(shape, name)
    style_fill(shape, fill)
    if line is None:
        shape.line.fill.background()
    else:
        style_line(shape, line, line_width)
    return shape


def set_text_frame(
    shape,
    lines: Sequence[str],
    *,
    font_size: float,
    color: str,
    bold: bool = False,
    font_name: str = FONT_UI,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    line_spacing: float = 1.0,
    space_after: float = 0.0,
) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)

    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    font_size: float = 18.0,
    color: str = COLORS["ink"],
    bold: bool = False,
    font_name: str = FONT_UI,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    name: str = "body:text",
    line_spacing: float = 1.0,
    space_after: float = 0.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_name(shape, name)
    set_text_frame(
        shape,
        text.splitlines() or [""],
        font_size=font_size,
        color=color,
        bold=bold,
        font_name=font_name,
        align=align,
        valign=valign,
        margin=margin,
        line_spacing=line_spacing,
        space_after=space_after,
    )
    return shape


def add_box_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill: str,
    line: str,
    font_size: float = 16.0,
    color: str = COLORS["ink"],
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
    name: str = "label:box",
    margin: float = 0.08,
):
    shape = add_rect(slide, x, y, w, h, fill=fill, line=line, name=name)
    set_text_frame(
        shape,
        text.splitlines(),
        font_size=font_size,
        color=color,
        bold=bold,
        align=align,
        valign=MSO_ANCHOR.MIDDLE,
        margin=margin,
        line_spacing=0.95,
    )
    return shape


def add_slide_frame(slide, number: int, stage: str, stage_color: str) -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(COLORS["paper"])

    add_rect(slide, 0, 0, 0.12, SLIDE_H, fill=stage_color, name="decor:accent")
    add_text(
        slide,
        CONTENT_X,
        0.20,
        4.0,
        0.22,
        stage,
        font_size=10.0,
        color=stage_color,
        bold=True,
        name="small:stage",
    )
    add_rect(
        slide,
        CONTENT_X,
        7.04,
        CONTENT_W,
        0.012,
        fill=COLORS["line"],
        name="decor:footer-line",
    )
    add_text(
        slide,
        CONTENT_X,
        7.14,
        5.2,
        0.16,
        "DISPLAY_ON 数据上行链路",
        font_size=9.0,
        color=COLORS["muted"],
        name="footer:name",
    )
    add_text(
        slide,
        11.48,
        7.14,
        1.22,
        0.16,
        f"{number:02d} / 12",
        font_size=9.0,
        color=COLORS["muted"],
        align=PP_ALIGN.RIGHT,
        name="footer:number",
    )


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(
        slide,
        CONTENT_X,
        0.52,
        CONTENT_W,
        0.48,
        title,
        font_size=30.0,
        color=COLORS["ink"],
        bold=True,
        name="title:main",
    )
    if subtitle:
        add_text(
            slide,
            CONTENT_X,
            1.02,
            CONTENT_W,
            0.30,
            subtitle,
            font_size=15.0,
            color=COLORS["muted"],
            name="label:subtitle",
        )


def add_bullets(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    bullets: Iterable[str],
    *,
    color: str = COLORS["ink"],
    font_size: float = 18.0,
    name: str = "body:bullets",
    space_after: float = 8.0,
) -> None:
    lines = [f"• {item}" for item in bullets]
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_name(shape, name)
    set_text_frame(
        shape,
        lines,
        font_size=font_size,
        color=color,
        font_name=FONT_UI,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
        margin=0.02,
        line_spacing=1.05,
        space_after=space_after,
    )


def add_code_block(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    path: str,
    code: str,
    *,
    code_size: float = 14.0,
    name: str = "code:excerpt",
) -> None:
    add_rect(slide, x, y, w, h, fill=COLORS["code"], name="decor:code-bg")
    add_text(
        slide,
        x + 0.16,
        y + 0.10,
        w - 0.32,
        0.20,
        path,
        font_size=9.0,
        color=COLORS["code_muted"],
        name="small:code-path",
    )
    shape = slide.shapes.add_textbox(
        Inches(x + 0.16),
        Inches(y + 0.37),
        Inches(w - 0.32),
        Inches(h - 0.47),
    )
    set_shape_name(shape, name)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, line in enumerate(code.strip("\n").splitlines()):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 0.90
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT_CODE
        run.font.size = Pt(code_size)
        run.font.color.rgb = rgb(COLORS["white"])


def add_arrow(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    *,
    color: str,
    reverse: bool = False,
    name: str = "label:arrow",
) -> None:
    shape_type = MSO_SHAPE.LEFT_ARROW if reverse else MSO_SHAPE.RIGHT_ARROW
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    set_shape_name(shape, name)
    style_fill(shape, color)
    shape.line.fill.background()
    set_text_frame(
        shape,
        [label],
        font_size=11.0,
        color=COLORS["white"],
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0.02,
    )


def add_io_summary(
    slide,
    y: float,
    input_text: str,
    process_text: str,
    output_text: str,
) -> None:
    add_rect(
        slide,
        CONTENT_X,
        y,
        CONTENT_W,
        0.62,
        fill=COLORS["soft"],
        line=COLORS["line"],
        name="decor:io-band",
    )
    cell_w = CONTENT_W / 3
    for index in (1, 2):
        add_rect(
            slide,
            CONTENT_X + cell_w * index,
            y + 0.09,
            0.012,
            0.44,
            fill=COLORS["line"],
            name="decor:io-separator",
        )
    items = (("输入", input_text), ("处理", process_text), ("输出", output_text))
    for index, (label, value) in enumerate(items):
        x = CONTENT_X + cell_w * index + 0.15
        add_text(
            slide,
            x,
            y + 0.10,
            0.58,
            0.18,
            label,
            font_size=10.0,
            color=COLORS["muted"],
            bold=True,
            name="small:io-label",
        )
        add_text(
            slide,
            x + 0.62,
            y + 0.08,
            cell_w - 0.82,
            0.38,
            value,
            font_size=14.0,
            color=COLORS["ink"],
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
            name="label:io-value",
        )


def add_speaker_cue(slide, text: str) -> None:
    add_text(
        slide,
        CONTENT_X,
        6.72,
        CONTENT_W,
        0.20,
        f"讲解重点：{text}",
        font_size=10.0,
        color=COLORS["muted"],
        name="small:speaker-cue",
    )


def add_step_node(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    number: str,
    title: str,
    detail: str,
    *,
    color: str,
    fill: str,
) -> None:
    add_rect(slide, x, y, w, h, fill=fill, line=color, line_width=1.2, name="decor:step")
    add_box_text(
        slide,
        x + 0.12,
        y + 0.14,
        0.42,
        0.36,
        number,
        fill=color,
        line=color,
        font_size=12.0,
        color=COLORS["white"],
        bold=True,
        name="label:step-number",
        margin=0,
    )
    add_text(
        slide,
        x + 0.66,
        y + 0.12,
        w - 0.78,
        0.30,
        title,
        font_size=16.0,
        color=COLORS["ink"],
        bold=True,
        name="label:step-title",
    )
    add_text(
        slide,
        x + 0.16,
        y + 0.57,
        w - 0.32,
        h - 0.67,
        detail,
        font_size=14.0,
        color=COLORS["muted"],
        name="label:step-detail",
        line_spacing=1.0,
    )


def new_slide(prs: Presentation, number: int, stage: str, stage_color: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_frame(slide, number, stage, stage_color)
    return slide


def build_slide_1(prs: Presentation) -> None:
    slide = new_slide(prs, 1, "全景", COLORS["request"])
    add_text(
        slide,
        0.86,
        1.18,
        11.7,
        0.68,
        "一条数据的旅程",
        font_size=42.0,
        color=COLORS["ink"],
        bold=True,
        align=PP_ALIGN.CENTER,
        name="title:cover",
    )
    add_text(
        slide,
        0.86,
        1.96,
        11.7,
        0.46,
        "从前端请求到 FSM 返回，再到页面显示",
        font_size=24.0,
        color=COLORS["muted"],
        align=PP_ALIGN.CENTER,
        name="body:cover-subtitle",
    )
    add_text(
        slide,
        0.86,
        2.54,
        11.7,
        0.28,
        "ArkTS 前端 × Go 后端 × FSM  |  30 分钟代码链路培训",
        font_size=14.0,
        color=COLORS["request"],
        bold=True,
        align=PP_ALIGN.CENTER,
        name="label:cover-meta",
    )

    steps = ("连接", "请求", "路由", "指令", "返回", "解析", "转状态", "显示")
    x = 0.72
    for index, step in enumerate(steps):
        color = COLORS["request"] if index < 4 else COLORS["response"]
        fill = COLORS["request_light"] if index < 4 else COLORS["response_light"]
        add_box_text(
            slide,
            x,
            3.62,
            1.30,
            0.72,
            f"{index + 1:02d}\n{step}",
            fill=fill,
            line=color,
            font_size=15.0,
            color=color,
            bold=True,
            name="label:cover-step",
        )
        if index < len(steps) - 1:
            add_arrow(
                slide,
                x + 1.31,
                3.86,
                0.18,
                0.22,
                "",
                color=COLORS["line"],
                name="label:cover-arrow",
            )
        x += 1.50

    add_box_text(
        slide,
        2.25,
        5.12,
        8.83,
        0.76,
        "先看懂完整链路，再看每一段对应代码",
        fill=COLORS["white"],
        line=COLORS["line"],
        font_size=20.0,
        color=COLORS["ink"],
        bold=True,
        name="body:cover-message",
    )


def build_slide_2(prs: Presentation) -> None:
    slide = new_slide(prs, 2, "角色与边界", COLORS["request"])
    add_title(slide, "三个角色，一条链路", "先明确谁做什么，后面的代码才不会串层")

    columns = (
        (
            0.66,
        "ArkTS 前端",
            "WebSocket JSON",
            ("连接 127.0.0.1:18080", "发送请求 JSON", "接收广播数据帧", "映射并写入页面状态"),
            COLORS["backend"],
            COLORS["backend_light"],
        ),
        (
            4.64,
            "Go 后端",
            "协议翻译 + 数据整形",
            ("路由前端请求", "向 FSM 发送 SYNC 指令", "监听 1128 并解析", "缓存、聚合、广播"),
            COLORS["backend"],
            COLORS["backend_light"],
        ),
        (
            8.62,
            "FSM",
            "CTCP SYNC 二进制",
            ("1279 接收 DISPLAY_ON", "主案例回连 HC:1128", "返回配置 StGlobal", "周期推送 Statistics"),
            COLORS["backend"],
            COLORS["backend_light"],
        ),
    )

    for x, title, protocol, bullets, color, fill in columns:
        add_rect(slide, x, 1.50, 3.76, 4.28, fill=COLORS["white"], line=COLORS["line"], name="decor:role")
        add_rect(slide, x, 1.50, 3.76, 0.62, fill=color, name="decor:role-header")
        add_text(
            slide,
            x + 0.18,
            1.65,
            3.40,
            0.30,
            title,
            font_size=20.0,
            color=COLORS["white"],
            bold=True,
            name="body:role-title",
        )
        add_box_text(
            slide,
            x + 0.20,
            2.35,
            3.36,
            0.54,
            protocol,
            fill=fill,
            line=color,
            font_size=15.0,
            color=color,
            bold=True,
            name="label:role-protocol",
        )
        add_bullets(slide, x + 0.22, 3.12, 3.30, 2.25, bullets, font_size=18.0, name="body:role-bullets")

    add_box_text(
        slide,
        2.18,
        6.02,
        8.98,
        0.52,
        "术语：FSM = 设备侧控制程序  |  SYNC = CTCP 固定帧头  |  topic = 数据类别字段",
        fill=COLORS["soft"],
        line=COLORS["line"],
        font_size=15.0,
        color=COLORS["ink"],
        bold=True,
        name="label:role-summary",
    )


def build_slide_3(prs: Presentation) -> None:
    slide = new_slide(prs, 3, "完整链路", COLORS["request"])
    add_title(slide, "完整链路：连接、请求、指令、返回、显示", "请求向右走，返回向左走；FSM 返回不是同步函数返回")

    actor_specs = (
        (0.72, "ArkTS 前端", COLORS["backend"]),
        (5.02, "Go 后端", COLORS["backend"]),
        (9.32, "FSM", COLORS["backend"]),
    )
    for x, actor, color in actor_specs:
        add_box_text(
            slide,
            x,
            1.40,
            3.28,
            0.52,
            actor,
            fill=COLORS["white"],
            line=color,
            font_size=18.0,
            color=color,
            bold=True,
            name="body:actor",
        )

    add_step_node(
        slide,
        0.72,
        2.18,
        3.28,
        1.42,
        "1",
        "WebSocket open",
        "4× requestStGlobal\n+ requestHomeStats\n+ requestStatistics",
        color=COLORS["request"],
        fill=COLORS["request_light"],
    )
    add_arrow(slide, 4.10, 2.62, 0.82, 0.48, "JSON", color=COLORS["request"])
    add_step_node(
        slide,
        5.02,
        2.18,
        3.28,
        1.42,
        "2",
        "handleIncoming",
        "仅 requestStGlobal\n触发 DISPLAY_ON",
        color=COLORS["request"],
        fill=COLORS["request_light"],
    )
    add_arrow(slide, 8.40, 2.62, 0.82, 0.48, "1279", color=COLORS["request"])
    add_step_node(
        slide,
        9.32,
        2.18,
        3.28,
        1.42,
        "3",
        "接收 DISPLAY_ON",
        "16 字节 SYNC 帧\ncmd = 0x0019",
        color=COLORS["request"],
        fill=COLORS["request_light"],
    )

    add_step_node(
        slide,
        9.32,
        4.18,
        3.28,
        1.48,
        "4",
        "另起连接返回",
        "FSM -> HC:1128\nConfig / Statistics",
        color=COLORS["response"],
        fill=COLORS["response_light"],
    )
    add_arrow(slide, 8.40, 4.65, 0.82, 0.48, "SYNC", color=COLORS["response"], reverse=True)
    add_step_node(
        slide,
        5.02,
        4.18,
        3.28,
        1.48,
        "5",
        "读取、解析、缓存",
        "配置即时广播\n统计每 1 秒取快照",
        color=COLORS["response"],
        fill=COLORS["response_light"],
    )
    add_arrow(slide, 4.10, 4.65, 0.82, 0.48, "JSON", color=COLORS["response"], reverse=True)
    add_step_node(
        slide,
        0.72,
        4.18,
        3.28,
        1.48,
        "6",
        "分流并显示",
        "数据类别字段 -> 对应状态分支\n配置管理器 / KPI 状态 -> UI",
        color=COLORS["response"],
        fill=COLORS["response_light"],
    )

    add_box_text(
        slide,
        3.45,
        5.92,
        6.44,
        0.52,
        "关键：FSM 是反向新建连接，不是原 1279 连接上的同步 return",
        fill=COLORS["warning_light"],
        line=COLORS["warning"],
        font_size=15.0,
        color=COLORS["warning"],
        bold=True,
        name="label:async-return",
    )
    add_speaker_cue(slide, "先让大家记住方向和边界，后面每页只放大一个节点。")


def add_technical_panel(
    slide,
    heading: str,
    bullets: Sequence[str],
    *,
    accent: str,
    fill: str,
    conclusion: str,
) -> None:
    add_rect(
        slide,
        8.18,
        1.48,
        4.52,
        4.02,
        fill=COLORS["white"],
        line=COLORS["line"],
        name="decor:technical-panel",
    )
    add_rect(
        slide,
        8.18,
        1.48,
        0.08,
        3.92,
        fill=accent,
        name="decor:technical-accent",
    )
    add_text(
        slide,
        8.48,
        1.72,
        3.92,
        0.34,
        heading,
        font_size=20.0,
        color=COLORS["ink"],
        bold=True,
        name="body:technical-heading",
    )
    add_bullets(
        slide,
        8.46,
        2.16,
        3.92,
        2.48,
        bullets,
        font_size=17.0,
        name="label:technical-bullets",
        space_after=4.0,
    )
    add_box_text(
        slide,
        8.46,
        4.84,
        3.94,
        0.50,
        conclusion,
        fill=fill,
        line=accent,
        font_size=14.0,
        color=accent,
        bold=True,
        name="label:technical-conclusion",
    )


def build_slide_4(prs: Presentation) -> None:
    slide = new_slide(prs, 4, "前端连接", COLORS["request"])
    add_title(
        slide,
        "连接成功，只代表通道已经建立",
        "open 回调会立刻发起首轮请求，但此时还没有任何 FSM 业务数据",
    )
    frontend_client = FRONTEND_ROOT / "entry/src/main/ets/utils/network/HarmonyWebSocketClient.ets"
    code = source_excerpt(frontend_client, ((689, 689), (692, 697)))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/new/my_harmony/entry/src/main/ets/utils/network/HarmonyWebSocketClient.ets | openCallback",
        code,
    )
    add_technical_panel(
        slide,
        "这里发生了什么",
        (
            "state = open；不等待 ready",
            "循环请求 4 个 FSM 的 StGlobal",
            "再请求 homeStats / statistics 快照",
            "失败、重连、业务回包都是后续状态",
        ),
        accent=COLORS["request"],
        fill=COLORS["request_light"],
        conclusion="open != ready != 业务数据已返回",
    )
    add_io_summary(
        slide,
        5.66,
        "WebSocket open 回调",
        "设置状态并发送 6 个 JSON 请求",
        "请求已发出，尚未拿到 FSM 数据",
    )
    add_speaker_cue(slide, "把“通道可用”和“业务数据到达”分开，排查时先看是哪一步没发生。")


def build_slide_5(prs: Presentation) -> None:
    slide = new_slide(prs, 5, "前端请求", COLORS["request"])
    add_title(
        slide,
        "前端发的是 JSON，不是 SYNC",
        "前端只描述“想做什么”，Go 再负责把请求翻译成设备协议",
    )
    frontend_client = FRONTEND_ROOT / "entry/src/main/ets/utils/network/HarmonyWebSocketClient.ets"
    code = source_excerpt(frontend_client, ((793, 804),))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/new/my_harmony/entry/src/main/ets/utils/network/HarmonyWebSocketClient.ets | requestStGlobal",
        code,
    )
    add_technical_panel(
        slide,
        "三类请求不是一回事",
        (
            "requestStGlobal 带 fsmId，拉设备配置",
            "requestHomeStats 尝试推聚合缓存",
            "requestStatistics 尝试推统计缓存",
            "只有 requestStGlobal 下发 DISPLAY_ON",
        ),
        accent=COLORS["request"],
        fill=COLORS["request_light"],
        conclusion="JSON：type=requestStGlobal, fsmId=256",
    )
    add_io_summary(
        slide,
        5.66,
        "type + 可选 fsmId",
        "JSON.stringify 后走 WebSocket 文本帧",
        "Go 收到控制消息",
    )
    add_speaker_cue(slide, "前端不知道 1279、SYNC 或 0x0019，这些都属于 Go 与 FSM 的边界。")


def build_slide_6(prs: Presentation) -> None:
    slide = new_slide(prs, 6, "后端路由", COLORS["request"])
    add_title(
        slide,
        "Go 按 type 路由请求",
        "handleIncoming 读取 type；requestStGlobal 进入 handler 后先回缓存、再异步访问 FSM",
    )
    websocket_go = GO_ROOT / "internal/runtime/websocket.go"
    handlers_go = GO_ROOT / "internal/runtime/websocket_handlers.go"
    code = source_excerpt(websocket_go, ((411, 412),))
    code += "\n...\n" + source_excerpt(handlers_go, ((12, 12), (14, 18), (21, 25)))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/goTest/go/ohos/Tcp/internal/runtime/websocket.go + websocket_handlers.go | handler",
        code,
    )
    add_technical_panel(
        slide,
        "路由后的差异",
        (
            "每条 request 只进入一次 handler",
            "5 个 sendLatest... 是本地缓存回放",
            "go func 只启动 1 个异步任务",
            "设备指令在 requestStGlobalFromFSM 内发送",
        ),
        accent=COLORS["backend"],
        fill=COLORS["backend_light"],
        conclusion="一次 handler = 先回缓存，再发 1 次 DISPLAY_ON",
    )
    add_io_summary(
        slide,
        5.66,
        "webSocketControlMessage",
        "按 type 路由；缓存回放；goroutine",
        "5 类缓存帧 + 1 次异步 FSM 请求",
    )
    add_speaker_cue(slide, "这里是协议边界入口：同样是前端请求，只有一条分支继续走设备网络。")


def build_slide_7(prs: Presentation) -> None:
    slide = new_slide(prs, 7, "发送指令", COLORS["request"])
    add_title(
        slide,
        "Go 把 JSON 翻译成 DISPLAY_ON",
        "Go 连接 FSM 的 1279 端口，发送 16 字节小端 SYNC 头；payload = 命令后的业务数据体",
    )
    ctcp_client = GO_ROOT / "client/ctcp_client.go"
    code = source_excerpt(ctcp_client, ((465, 469), (337, 343)))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/goTest/go/ohos/Tcp/client/ctcp_client.go | RequestStGlobalFromFSM / SendCMD.Bytes",
        code,
    )

    add_text(
        slide,
        8.30,
        1.52,
        4.18,
        0.30,
        "16 字节头 = 4 个 uint32",
        font_size=20.0,
        color=COLORS["ink"],
        bold=True,
        name="body:sync-heading",
    )
    fields = (
        ("0-3", "SYNC", "0x434E5953"),
        ("4-7", "src", "0x1000 / HC"),
        ("8-11", "dest", "0x0100... / FSM"),
        ("12-15", "cmd", "0x0019 / DISPLAY_ON"),
    )
    y = 2.02
    for byte_range, label, value in fields:
        add_box_text(
            slide,
            8.30,
            y,
            0.82,
            0.58,
            byte_range,
            fill=COLORS["request"],
            line=COLORS["request"],
            font_size=11.0,
            color=COLORS["white"],
            bold=True,
            name="label:sync-range",
        )
        add_box_text(
            slide,
            9.12,
            y,
            1.10,
            0.58,
            label,
            fill=COLORS["request_light"],
            line=COLORS["request"],
            font_size=13.0,
            color=COLORS["request"],
            bold=True,
            name="label:sync-field",
        )
        add_box_text(
            slide,
            10.22,
            y,
            2.28,
            0.58,
            value,
            fill=COLORS["white"],
            line=COLORS["line"],
            font_size=13.0,
            color=COLORS["ink"],
            bold=True,
            name="label:sync-value",
        )
        y += 0.70
    add_box_text(
        slide,
        8.30,
        4.94,
        4.20,
        0.46,
        "目标端口 1279  |  payload = 0 字节",
        fill=COLORS["warning_light"],
        line=COLORS["warning"],
        font_size=14.0,
        color=COLORS["warning"],
        bold=True,
        name="label:sync-target",
    )
    add_io_summary(
        slide,
        5.66,
        "fsmId + DISPLAY_ON",
        "LittleEndian 写入 16 字节 SYNC 头",
        "TCP 发往 FSM:1279",
    )
    add_speaker_cue(slide, "JSON 到这里结束；从这一页开始，链路变成设备二进制协议。")


def build_slide_8(prs: Presentation) -> None:
    slide = new_slide(prs, 8, "FSM 返回", COLORS["response"])
    add_title(
        slide,
        "FSM 从 1128 返回，Go 分三段读取",
        "FSM 新建反向 TCP 连接；这不是 1279 连接上的同步返回值",
    )
    ctcp_server = GO_ROOT / "internal/runtime/ctcp_server.go"
    code = source_excerpt(ctcp_server, ((306, 306), (311, 316), (328, 328), (349, 349)))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/goTest/go/ohos/Tcp/internal/runtime/ctcp_server.go | handleConnection",
        code,
    )
    add_technical_panel(
        slide,
        "三段读取",
        (
            "4 字节 SYNC：必须是 ASCII 'SYNC'",
            "12 字节头：src / dst / cmd",
            "payload：read-until-idle（500ms 或 EOF）",
            "dst 不是 HC_ID 直接丢弃",
        ),
        accent=COLORS["response"],
        fill=COLORS["response_light"],
        conclusion="主案例监听 0.0.0.0:1128，连接方向 FSM -> Go",
    )
    add_io_summary(
        slide,
        5.66,
        "FSM 新建 TCP 连接",
        "SYNC -> command head -> payload",
        "交给 handleCommandPayload",
    )
    add_speaker_cue(slide, "端口、连接方向和读取顺序是抓包或看日志时最重要的三个坐标。")


def build_slide_9(prs: Presentation) -> None:
    slide = new_slide(prs, 9, "解析与缓存", COLORS["response"])
    add_title(
        slide,
        "配置即时转发，统计先缓存再定时广播",
        "同一个 1128 入口，根据 cmd 走两条不同的处理分支",
    )
    ctcp_server = GO_ROOT / "internal/runtime/ctcp_server.go"
    config_code = source_excerpt(
        ctcp_server,
        ((355, 357), (369, 369), (384, 384), (391, 392), (397, 400)),
    )
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/goTest/go/ohos/Tcp/internal/runtime/ctcp_server.go | handleCommandPayload",
        config_code,
    )
    add_technical_panel(
        slide,
        "两类命令，两种时机",
        (
            "0x1000 -> StGlobal -> 立即 JSON",
            "0x1001 -> StStatistics -> 最新快照",
            "ParseData 按 Go 结构体大小读取",
            "Go 布局必须与 FSM 二进制一致",
        ),
        accent=COLORS["warning"],
        fill=COLORS["warning_light"],
        conclusion="缓存是节流点，不是另一套业务数据源",
    )
    add_io_summary(
        slide,
        5.66,
        "cmd + 二进制 payload",
        "解析结构体并按命令分支",
        "stglobal 立即广播 / statistics 写缓存",
    )
    add_speaker_cue(slide, "不要把“FSM 已返回”和“前端已收到”画成同一个时刻，统计中间还有缓存与定时器。")


def build_slide_10(prs: Presentation) -> None:
    slide = new_slide(prs, 10, "后端广播", COLORS["response"])
    add_title(
        slide,
        "Go 每秒取快照，再广播给所有连接",
        "topic 是帧内的数据类别字段；hub 不按类别维护客户端名单",
    )
    websocket_go = GO_ROOT / "internal/runtime/websocket.go"
    publish_code = source_excerpt(websocket_go, ((642, 647), (317, 324)))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        3.92,
        "E:/goTest/go/ohos/Tcp/internal/runtime/websocket.go | data frame + hub.run",
        publish_code,
    )
    add_technical_panel(
        slide,
        "广播前后的规则",
        (
            "定时器每 1 秒；请求也可立即触发",
            "超过 2.5 秒未更新，显示速度归零",
            "封装 type/topic/data/at，topic 转小写",
            "hub 投递给每个已连接客户端",
        ),
        accent=COLORS["response"],
        fill=COLORS["response_light"],
        conclusion="topic 负责标识数据，前端负责按字段分流",
    )
    add_io_summary(
        slide,
        5.66,
        "最新 StStatistics 缓存",
        "1 秒快照 + JSON data 帧 + hub 广播",
        "statistics 与 homeStats 数据帧",
    )
    add_speaker_cue(slide, "Publish 是函数名；从 hub 的 clients 循环可以直接看出实际行为是全连接广播。")


def build_slide_11(prs: Presentation) -> None:
    slide = new_slide(prs, 11, "前端映射与显示", COLORS["response"])
    add_title(
        slide,
        "前端按数据类别字段分流，状态变化驱动显示",
        "映射 = JSON 字段转前端结构/状态；statistics 与 homeStats 走不同分支",
    )
    frontend_client = FRONTEND_ROOT / "entry/src/main/ets/utils/network/HarmonyWebSocketClient.ets"
    sorting_card = FRONTEND_ROOT / "entry/src/main/ets/pages/home/SortingInfoCard.ets"
    route_code = source_excerpt(frontend_client, ((1728, 1734),))
    state_code = source_excerpt(frontend_client, ((2155, 2155), (2182, 2182), (2708, 2708)))
    state_code += "\n...\n" + source_excerpt(sorting_card, ((20, 20),))
    add_code_block(
        slide,
        0.62,
        1.48,
        7.30,
        1.82,
        "E:/new/my_harmony/entry/src/main/ets/utils/network/HarmonyWebSocketClient.ets",
        route_code,
    )
    add_code_block(
        slide,
        0.62,
        3.42,
        7.30,
        1.98,
        "E:/new/my_harmony | HarmonyWebSocketClient.ets + pages/home/SortingInfoCard.ets",
        state_code,
    )

    add_text(
        slide,
        8.28,
        1.58,
        4.18,
        0.34,
        "两个真实状态分支",
        font_size=20.0,
        color=COLORS["ink"],
        bold=True,
        name="body:frontend-branches",
    )
    branch_specs = (
        (2.10, "statistics", "Mapper", "GlobalData", COLORS["response"], COLORS["response_light"]),
        (3.34, "homeStats", "AppStorage", "页面绑定", COLORS["response"], COLORS["response_light"]),
    )
    for y, source, middle, target, color, fill in branch_specs:
        add_box_text(
            slide,
            8.28,
            y,
            1.12,
            0.64,
            source,
            fill=fill,
            line=color,
            font_size=12.0,
            color=color,
            bold=True,
            name="label:state-source",
        )
        add_arrow(slide, 9.48, y + 0.12, 0.42, 0.40, "", color=color)
        add_box_text(
            slide,
            9.92,
            y,
            1.16,
            0.64,
            middle,
            fill=COLORS["white"],
            line=COLORS["line"],
            font_size=12.0,
            color=COLORS["ink"],
            bold=True,
            name="label:state-middle",
        )
        add_arrow(slide, 11.14, y + 0.12, 0.42, 0.40, "", color=color)
        add_box_text(
            slide,
            11.56,
            y,
            1.10,
            0.64,
            target,
            fill=fill,
            line=color,
            font_size=12.0,
            color=color,
            bold=True,
            name="label:state-target",
        )
    add_box_text(
        slide,
        8.28,
        4.62,
        4.38,
        0.56,
        "stglobal 走第三条分支：映射后更新设备配置",
        fill=COLORS["warning_light"],
        line=COLORS["warning"],
        font_size=13.0,
        color=COLORS["warning"],
        bold=True,
        name="label:stglobal-branch",
    )
    add_io_summary(
        slide,
        5.66,
        "type=data + topic + data",
        "字段分流、结构映射、写状态",
        "GlobalDataInterface / AppStorage -> UI",
    )
    add_speaker_cue(slide, "前端没有按类别维护连接；它只是收到广播后读取 topic 字段，再进入对应处理函数。")


def build_slide_12(prs: Presentation) -> None:
    slide = new_slide(prs, 12, "排查与复盘", COLORS["warning"])
    add_title(
        slide,
        "八个节点定位：数据停在哪里",
        "按链路顺序排查，不要一上来只盯页面或只盯 FSM",
    )
    nodes = (
        ("1", "WebSocket open", "前端 state=open", COLORS["request"], COLORS["request_light"]),
        ("2", "JSON 请求", "sendText=true", COLORS["request"], COLORS["request_light"]),
        ("3", "Go 路由", "handleIncoming", COLORS["request"], COLORS["request_light"]),
        ("4", "指令发出", "DISPLAY_ON -> 1279", COLORS["request"], COLORS["request_light"]),
        ("5", "FSM 回连", "FSM -> 1128", COLORS["response"], COLORS["response_light"]),
        ("6", "解析缓存", "0x1000 / 0x1001", COLORS["response"], COLORS["response_light"]),
        ("7", "WebSocket 广播", "type/topic/data", COLORS["response"], COLORS["response_light"]),
        ("8", "映射显示", "状态 -> UI", COLORS["response"], COLORS["response_light"]),
    )
    for index, (number, title, detail, color, fill) in enumerate(nodes):
        row = index // 4
        col = index % 4
        x = 0.68 + col * 3.12
        y = 1.62 + row * 1.74
        add_step_node(
            slide,
            x,
            y,
            2.76,
            1.30,
            number,
            title,
            detail,
            color=color,
            fill=fill,
        )
        if col < 3:
            add_arrow(
                slide,
                x + 2.80,
                y + 0.46,
                0.26,
                0.34,
                "",
                color=color,
                name="label:diagnostic-arrow",
            )

    add_box_text(
        slide,
        1.18,
        5.16,
        11.00,
        0.68,
        "复盘：连接 -> 请求 -> 路由 -> 指令 -> 返回 -> 解析 -> 广播 -> 显示",
        fill=COLORS["white"],
        line=COLORS["line"],
        font_size=19.0,
        color=COLORS["ink"],
        bold=True,
        name="body:recap",
    )
    add_box_text(
        slide,
        2.20,
        6.02,
        8.96,
        0.48,
        "先确认上一节点的输出，再检查下一节点的输入",
        fill=COLORS["warning_light"],
        line=COLORS["warning"],
        font_size=16.0,
        color=COLORS["warning"],
        bold=True,
        name="label:diagnostic-rule",
    )
    add_speaker_cue(slide, "主线先查 1128；1127 是图像口，属于另一条数据入口。缓存和 AppStorage 都不是新增数据源。")


def build_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    builders = (
        build_slide_1,
        build_slide_2,
        build_slide_3,
        build_slide_4,
        build_slide_5,
        build_slide_6,
        build_slide_7,
        build_slide_8,
        build_slide_9,
        build_slide_10,
        build_slide_11,
        build_slide_12,
    )
    for builder in builders:
        builder(prs)
    return prs


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation = build_deck()
    temporary = OUTPUT.with_name(OUTPUT.name + ".tmp")
    try:
        presentation.save(temporary)
        with ZipFile(temporary) as archive:
            if archive.testzip() is not None:
                raise RuntimeError("generated PPTX contains a corrupt member")
        if len(Presentation(temporary).slides) != 12:
            raise RuntimeError("generated PPTX does not contain 12 slides")
        os.replace(temporary, OUTPUT)
    finally:
        temporary.unlink(missing_ok=True)
    print(f"{len(presentation.slides)} slides written: {OUTPUT}")


if __name__ == "__main__":
    main()
